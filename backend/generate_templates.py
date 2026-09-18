"""一次性生成 Excel 和 Word 题目导入模板，放到 frontend/public/templates/。"""
import os

from docx import Document
from openpyxl import Workbook

OUT_DIR = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "..", "frontend", "public", "templates",
)
os.makedirs(OUT_DIR, exist_ok=True)

# ---------- Excel ----------
wb = Workbook()
ws = wb.active
ws.title = "题目"

ws.append(["题干", "选项A", "选项B", "选项C", "选项D", "正确答案", "解析"])
ws.append(["1 + 1 等于几？", "1", "2", "3", "4", "B", "1 + 1 = 2"])
ws.append(["中国的首都是？", "上海", "北京", "广州", "深圳", "B", "北京是中国的首都"])

# 列宽
for col, width in zip("ABCDEFG", [32, 12, 12, 12, 12, 10, 32]):
    ws.column_dimensions[col].width = width

wb.save(os.path.join(OUT_DIR, "题目导入模板.xlsx"))

# ---------- Word ----------
doc = Document()
doc.add_heading("题目导入模板", level=1)
doc.add_paragraph("请按以下格式填写，每道题之间空一行。正确答案支持 A/B/C/D 或 1/2/3/4。")
doc.add_paragraph("")

doc.add_paragraph("1. 1 + 1 等于几？")
doc.add_paragraph("A. 1")
doc.add_paragraph("B. 2")
doc.add_paragraph("C. 3")
doc.add_paragraph("D. 4")
doc.add_paragraph("答案：B")
doc.add_paragraph("解析：1 + 1 = 2")
doc.add_paragraph("")

doc.add_paragraph("2. 中国的首都是？")
doc.add_paragraph("A. 上海")
doc.add_paragraph("B. 北京")
doc.add_paragraph("C. 广州")
doc.add_paragraph("D. 深圳")
doc.add_paragraph("答案：B")
doc.add_paragraph("解析：北京是中国的首都")

doc.save(os.path.join(OUT_DIR, "题目导入模板.docx"))

print(f"模板已生成：{OUT_DIR}")
print("  - 题目导入模板.xlsx")
print("  - 题目导入模板.docx")