"""导入解析器单元测试。

覆盖两件事：

1. **v1 语义不能被破坏**：`parse_file()` 仍然是「遇到第一处错误就抛」，
   老调用方（如果还有）行为不变；
2. **v2 容错语义**：`parse_file_report()` 必须做到
   - 单行坏不影响其它行（行级隔离）
   - 每条错误带得上 location + reason（能定位到第几行）
   - 文件内重复题只保留一条，且记成 warning 而不是 error
   - 整文件级失败（格式不支持 / 空文件）仍然抛 ImportError
"""

import io
import json

import pytest
from docx import Document
from openpyxl import Workbook

from app.importer import (
    ImportError as ParseError,
    content_signature,
    detect_format,
    parse_excel,
    parse_excel_report,
    parse_file,
    parse_file_report,
    parse_json,
    parse_json_report,
    parse_pdf_report,
    parse_word_report,
)


def make_question(i: int, **overrides) -> dict:
    item = {
        "stem": f"第 {i} 题：下列说法正确的是？",
        "question_type": "single",
        "options": ["甲", "乙", "丙", "丁"],
        "correct_answer": "B",
        "explanation": "解析内容",
    }
    item.update(overrides)
    return item


def as_bytes(items) -> bytes:
    return json.dumps(items, ensure_ascii=False).encode("utf-8")


# ============================================================
# 格式识别
# ============================================================

@pytest.mark.parametrize("filename,expected", [
    ("a.json", "json"),
    ("A.XLSX", "excel"),
    ("b.xls", "excel"),
    ("c.docx", "word"),
    ("d.pdf", "pdf"),
])
def test_detect_format(filename, expected):
    assert detect_format(filename) == expected


def test_detect_format_rejects_unknown():
    with pytest.raises(ParseError):
        detect_format("notes.txt")


# ============================================================
# JSON
# ============================================================

def test_json_report_all_good():
    report = parse_json_report(as_bytes([make_question(i) for i in range(1, 6)]))
    assert report.success == 5
    assert report.failed == 0
    assert report.total_rows == 5
    assert report.success_rate == 1.0


def test_json_report_row_level_isolation():
    """一行坏，不能影响其它四行 —— 这是 v2 相对 v1 的核心差别。"""
    items = [make_question(i) for i in range(1, 6)]
    items[2]["correct_answer"] = ""          # 第 3 条：答案为空
    report = parse_json_report(as_bytes(items))

    assert report.success == 4
    assert report.failed == 1
    assert report.total_rows == 5
    issue = report.issues[0]
    assert issue.location == "第 3 条"
    assert issue.severity == "error"
    assert "答案" in issue.reason
    assert report.success_rate == 0.8


def test_json_report_collects_every_bad_row():
    """不能只报第一条错 —— 用户要一次看到所有要改的行。"""
    items = [make_question(i) for i in range(1, 11)]
    items[1]["stem"] = ""
    items[4]["options"] = ["只有一个"]
    items[7].pop("correct_answer")
    report = parse_json_report(as_bytes(items))

    assert report.success == 7
    assert report.failed == 3
    assert [i.location for i in report.issues] == ["第 2 条", "第 5 条", "第 8 条"]


def test_json_report_dedupes_with_warning():
    items = [make_question(1), make_question(1), make_question(2)]
    report = parse_json_report(as_bytes(items))

    assert report.success == 2
    assert report.deduped == 1
    assert report.failed == 0                      # 去重是 warning，不算失败
    assert report.issues[0].severity == "warning"
    assert "重复" in report.issues[0].reason


def test_json_v1_still_raises_on_first_error():
    """v1 兼容语义：parse_json 遇错即抛。"""
    items = [make_question(1), make_question(2, stem="")]
    with pytest.raises(ParseError) as exc:
        parse_json(as_bytes(items))
    assert "第 2 条" in str(exc.value)


def test_json_report_bad_syntax_is_file_level():
    with pytest.raises(ParseError):
        parse_json_report(b"{ not json")


def test_json_report_accepts_single_object():
    report = parse_json_report(as_bytes(make_question(1)))
    assert report.success == 1


# ============================================================
# Excel
# ============================================================

def build_xlsx(rows, header) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_excel_report_partial_success():
    header = ["题干", "题型", "选项A", "选项B", "正确答案", "解析"]
    rows = [
        ["好题一", "single", "甲", "乙", "A", "解析"],
        ["好题二", "single", "甲", "乙", "B", "解析"],
        ["坏题：答案是 E", "single", "甲", "乙", "E", "解析"],
        ["好题三", "judge", "", "", "A", ""],
    ]
    report = parse_excel_report(build_xlsx(rows, header))

    assert report.success == 3
    assert report.failed == 1
    assert report.issues[0].location == "第 4 行"     # 表头占第 1 行
    assert "超出选项范围" in report.issues[0].reason


def test_excel_option_columns_can_skip_letters():
    """回归：v1 遇到第一个缺失的选项列就 break。

    表头是「选项A / 选项B / 选项D」（手工整理时跳过 C 很常见），
    v1 会在 D 之前 break，D 选项直接丢失，答案 D 被判「超出选项范围」。

    v2 的做法是：丢掉空列后重新编号 A→A、B→B、D→C，
    并把答案里的字母按同一张映射表翻译过去，保证选项和答案不会错位。
    """
    header = ["题干", "题型", "选项A", "选项B", "选项D", "正确答案"]
    rows = [["跳列的题", "single", "甲", "乙", "丁", "D"]]
    report = parse_excel_report(build_xlsx(rows, header))

    assert report.success == 1, report.issues
    q = report.questions[0]
    assert q["options"] == ["甲", "乙", "丁"]      # D 列被正确收进来了
    assert q["correct_answer"] == "C"              # 答案 D 同步映射成 C
    assert any("重新编号" in i.reason for i in report.issues), "应该有重编号告警"


def test_excel_multi_letter_answer_is_remapped():
    """多选题同样要重映射：选项 A/C/D + 答案 "AD" → 选项 [甲,乙,丙] + 答案 "AC"。"""
    header = ["题干", "题型", "选项A", "选项C", "选项D", "正确答案"]
    rows = [["多选跳列", "multiple", "甲", "乙", "丙", "AD"]]
    report = parse_excel_report(build_xlsx(rows, header))

    assert report.success == 1, report.issues
    q = report.questions[0]
    assert q["options"] == ["甲", "乙", "丙"]
    assert q["correct_answer"] == "AC"


def test_excel_contiguous_columns_are_untouched():
    """选项列连续时不能触发重编号，答案是啥就是啥。"""
    header = ["题干", "题型", "选项A", "选项B", "选项C", "正确答案"]
    rows = [["正常题", "single", "甲", "乙", "丙", "C"]]
    report = parse_excel_report(build_xlsx(rows, header))

    assert report.success == 1, report.issues
    assert report.questions[0]["options"] == ["甲", "乙", "丙"]
    assert report.questions[0]["correct_answer"] == "C"
    assert not any("重新编号" in i.reason for i in report.issues)


def test_excel_missing_required_column_is_file_level():
    with pytest.raises(ParseError):
        parse_excel_report(build_xlsx([["只题干"]], ["题干"]))


def test_excel_empty_sheet_is_file_level():
    with pytest.raises(ParseError):
        parse_excel_report(build_xlsx([], ["题干", "正确答案"]))


def test_excel_v1_still_raises():
    header = ["题干", "正确答案", "选项A", "选项B"]
    rows = [["好的", "A", "甲", "乙"], ["坏的", "Z", "甲", "乙"]]
    with pytest.raises(ParseError):
        parse_excel(build_xlsx(rows, header))


# ============================================================
# Word / PDF 纯文本
# ============================================================

def build_docx(paragraphs) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


WORD_TEXT = [
    "1. 下列哪个是编程语言？",
    "A. Python",
    "B. HTML",
    "C. CSS",
    "答案：A",
    "解析：Python 是通用编程语言",
    "2. 简述 TCP 三次握手。",
    "答案：客户端 SYN，服务端 SYN+ACK，客户端 ACK",
    "3. 这道题没有写答案",
    "A. 甲",
    "B. 乙",
]


def test_word_report_skips_missing_answer_block():
    report = parse_word_report(build_docx(WORD_TEXT))

    # 第 1 题单选、第 2 题简答成功；第 3 题缺答案被跳过
    assert report.success == 2
    assert report.failed == 1
    assert report.issues[0].location == "第 3 题"
    assert "答案" in report.issues[0].reason
    types = [q["question_type"] for q in report.questions]
    assert types == ["single", "short"]


def test_word_v1_raises_on_missing_answer():
    with pytest.raises(ParseError):
        parse_file("t.docx", build_docx(WORD_TEXT))


def test_word_no_text_is_file_level():
    with pytest.raises(ParseError):
        parse_word_report(build_docx([""]))


def test_pdf_not_reachable_page_is_file_level():
    """给一段不是 PDF 的字节：不能抛裸异常，要转成 ImportError。"""
    with pytest.raises(ParseError):
        parse_pdf_report(b"this is definitely not a pdf")


# ============================================================
# 分发 & 报告结构
# ============================================================

def test_parse_file_report_dispatch():
    report = parse_file_report("q.json", as_bytes([make_question(1)]))
    assert report.success == 1


def test_parse_file_report_empty_file():
    with pytest.raises(ParseError):
        parse_file_report("q.json", b"")


def test_parse_file_v1_interface_unchanged():
    assert len(parse_file("q.json", as_bytes([make_question(1)]))) == 1


def test_report_to_dict_shape():
    report = parse_json_report(as_bytes([make_question(1), make_question(1)]))
    payload = report.to_dict()
    for key in ("total_rows", "imported", "failed", "deduped",
                "success_rate", "summary", "issues", "issues_truncated"):
        assert key in payload
    assert payload["deduped"] == 1


def test_report_issue_payload_is_json_serializable():
    report = parse_json_report(as_bytes([make_question(1, stem="")]))
    # 接口要把它塞进 HTTP 响应，必须能序列化
    json.dumps(report.to_dict(), ensure_ascii=False)


def test_content_signature_ignores_whitespace():
    a = {"stem": "1 + 1 等于几", "question_type": "single", "options": ["1", "2"]}
    b = {"stem": "1 +  1  等于几", "question_type": "single", "options": ["1", "2"]}
    assert content_signature(a) == content_signature(b)
